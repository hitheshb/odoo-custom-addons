import io

from openpyxl import Workbook, load_workbook

HEADERS = ['Name', 'Model', 'Group', 'Create', 'Read', 'Update', 'Delete', 'Domain', 'Kind']


def build_xlsx(rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Access Lines'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for column, width in zip('ABCDEFGHI', (30, 28, 34, 9, 9, 9, 9, 40, 14)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_xlsx(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    all_rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not all_rows:
        return [], []
    header = [str(cell).strip() if cell is not None else '' for cell in all_rows[0]]
    body = [(index, row) for index, row in enumerate(all_rows[1:], start=2)]
    return header, body
