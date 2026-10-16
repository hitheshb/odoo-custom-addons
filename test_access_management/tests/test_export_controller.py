import base64
import io

from openpyxl import load_workbook

from odoo.tests import HttpCase, tagged


@tagged('post_install', '-at_install')
class TestExportController(HttpCase):

    def test_export_profile_and_sample_upload(self):
        profile = self.env['test.access.profile'].create({'name': 'Sample'})
        from odoo.tools import file_open
        with file_open('test_access_management/static/sample/sample_access_profile.xlsx', 'rb') as sample:
            content = sample.read()
        wizard = self.env['test.access.profile.upload'].create({
            'profile_id': profile.id,
            'file': base64.b64encode(content).decode(),
        })
        wizard.action_upload()
        self.assertFalse(wizard.error_report)
        self.assertEqual(len(profile.line_ids), 3)

        self.authenticate('admin', 'admin')
        response = self.url_open(f'/test_access_management/export/profile/{profile.id}')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet['A1'].value, 'Name')

    def test_export_forbidden_for_non_admin(self):
        self.env['res.users'].create({
            'name': 'Plain', 'login': 'plain_user', 'password': 'plain_user',
            'group_ids': [(6, 0, [self.env.ref('base.group_user').id])],
        })
        self.authenticate('plain_user', 'plain_user')
        response = self.url_open('/test_access_management/export/profile/1')
        self.assertEqual(response.status_code, 403)
